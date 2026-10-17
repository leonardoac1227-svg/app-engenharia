# -*- coding: utf-8 -*-
"""Relatório de Lançamento por Categoria — planilha no modelo do escritório.

A equipe aparece só como presença: horas da semana por pessoa e, numa coluna
ao lado, as faltas de cada dia com a justificativa. Valor por funcionário não
entra — o pagamento da equipe é lançado no Caixa, como qualquer despesa.
"""

from __future__ import annotations

import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from ..utils import data_br

# ---------------------------------------------------------------- estilo

FONTE = "Calibri"
MOEDA = '"R$"\\ #,##0.00'
DATA_BR = "DD/MM/YYYY"

CINZA_CABECALHO = "FFBFBFBF"
CINZA_STATUS = "FFD8D8D8"
VERDE_PONTO = "FFDDE6D9"        # verde-acinzentado claro: seção de presença
AMARELO_PENDENTE = "FFFFF2CC"  # falta ainda sem justificativa

LARGURAS = {"A": 33.0, "B": 48.7, "C": 26.5, "D": 11.7, "E": 16.3, "F": 13.2}
LARGURA_OCORRENCIAS = 62.0

_borda_fina = Side(style="thin")
BORDA = Border(left=_borda_fina, right=_borda_fina, top=_borda_fina, bottom=_borda_fina)

CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
ESQUERDA = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _preencher(cor: str) -> PatternFill:
    return PatternFill(start_color=cor, end_color=cor, fill_type="solid")


def _safe_float(valor) -> float:
    """SALVA-VIDAS: Garante que a conversão de valores nunca quebre o servidor, mesmo se vier texto sujo ou nulo."""
    if valor is None or valor == "":
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    
    texto = str(valor).upper().replace("R$", "").strip()
    if "." in texto and "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".")
        
    try:
        return float(texto)
    except ValueError:
        return 0.0


def _sigla_da_obra(obra: dict) -> str:
    sigla = (obra.get("sigla") or "").strip()
    if sigla:
        return sigla.upper()[:12]

    nome = (obra.get("nome") or "OBRA").upper()
    candidatos = [p for p in re.split(r"[\s\-/]+", nome) if any(c.isdigit() for c in p)]
    if candidatos:
        return candidatos[-1][:12]

    iniciais = "".join(p[0] for p in nome.split() if p)[:4]
    return iniciais or "OBRA"


def _altura_da_linha(*textos: str) -> float:
    linhas = 1
    for texto, largura in zip(textos, (33, 48, 26, 11)):
        if not texto:
            continue
        conteudo = str(texto)
        quebras = conteudo.count("\n") + 1
        estimadas = max(quebras, -(-len(conteudo) // largura))
        linhas = max(linhas, estimadas)
    return min(18 + linhas * 16, 150)


# ---------------------------------------------------------------- montagem

class _Aba:
    def __init__(self, planilha, obra: dict, dados: dict, com_ocorrencias: bool = False):
        self.com_ocorrencias = com_ocorrencias
        self.ultima_coluna = "G" if com_ocorrencias else "F"
        self.ws = planilha
        self.obra = obra or {}
        self.dados = dados or {}
        self.sigla = _sigla_da_obra(self.obra)
        self.semana = self.dados.get("semana_numero", 0)
        self.linha = 1
        self.celulas_de_secao: list[str] = []

    def cabecalho(self) -> None:
        ws = self.ws
        for letra, largura in LARGURAS.items():
            ws.column_dimensions[letra].width = largura
        if self.com_ocorrencias:
            ws.column_dimensions["G"].width = LARGURA_OCORRENCIAS

        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.print_area = f"A1:{self.ultima_coluna}1"
        ws.print_title_rows = "1:5"
        ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4)
        ws.freeze_panes = "A6"
        
        # Protegido contra campos nulos
        nome_obra = (self.obra.get('nome') or 'OBRA DESCONHECIDA').upper()
        cliente = (self.obra.get('cliente') or '').upper()
        texto_obra = f"{nome_obra} - {cliente}" if cliente else nome_obra
        
        data_ini = data_br(self.dados.get('data_inicio') or '')
        data_fim = data_br(self.dados.get('data_fim') or '')

        titulos = [
            "RELATORIO DE LANÇAMENTO POR CATEGORIA",
            texto_obra,
            f"PAG  - {data_ini} A {data_fim}",
        ]
        for texto in titulos:
            ws.merge_cells(start_row=self.linha, start_column=1, end_row=self.linha, end_column=5)
            celula = ws.cell(self.linha, 1, texto)
            celula.font = Font(name=FONTE, size=15, bold=True)
            celula.fill = _preencher(CINZA_CABECALHO)
            celula.alignment = Alignment(horizontal="center", vertical="center")
            for coluna in range(1, 6):
                ws.cell(self.linha, coluna).border = BORDA
            ws.row_dimensions[self.linha].height = 21
            self.linha += 1

        ws.row_dimensions[self.linha].height = 15.75
        self.linha += 1

        for coluna, titulo in enumerate(["Fornecedor", "Descrição", "Data", "DOC", "Valor (R$)"], 1):
            celula = ws.cell(self.linha, coluna, titulo)
            celula.font = Font(name=FONTE, size=16, bold=True)
            celula.alignment = CENTRO
            celula.border = BORDA
        if self.com_ocorrencias:
            celula = ws.cell(self.linha, 7, "Ocorrências do ponto")
            celula.font = Font(name=FONTE, size=16, bold=True)
            celula.alignment = CENTRO
            celula.border = BORDA
        ws.row_dimensions[self.linha].height = 24.75
        self.linha += 1

    def secao_ponto(self, titulo: str, ponto: dict) -> None:
        """Presença da equipe: horas da semana e, ao lado, as faltas do dia a dia.

        Fica fora do total a pagar e não tem coluna de valor. Quem cumpriu a
        jornada todos os dias fica com a coluna de ocorrências em branco.
        """
        ws = self.ws
        linha = self.linha
        equipe = ponto.get("equipe") or []

        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=4)
        celula = ws.cell(linha, 1, titulo)
        celula.font = Font(name=FONTE, size=15, bold=True)
        celula.fill = _preencher(VERDE_PONTO)
        celula.alignment = Alignment(horizontal="center", vertical="center")
        total = ws.cell(linha, 5, f"{_horas_br(sum(p['horas'] for p in equipe))}")
        total.font = Font(name=FONTE, size=16, bold=True)
        total.fill = _preencher(VERDE_PONTO)
        total.alignment = CENTRO
        jornada = ws.cell(linha, 7, f"Jornada: {_horas_br(ponto.get('jornada') or 8)} de seg a sex")
        jornada.font = Font(name=FONTE, size=12, bold=True)
        jornada.fill = _preencher(VERDE_PONTO)
        jornada.alignment = CENTRO
        for coluna in (1, 2, 3, 4, 5, 7):
            ws.cell(linha, coluna).border = BORDA
        ws.row_dimensions[linha].height = 34.95
        self.linha += 1

        for pessoa in equipe:
            linha = self.linha
            dias = pessoa["dias_trabalhados"]
            valores = [
                pessoa["nome"],
                pessoa["funcao"] or "mão de obra",
                f"{dias} dia(s) trabalhado(s)",
                "PONTO",
                _horas_br(pessoa["horas"]),
            ]
            for coluna, valor in enumerate(valores, 1):
                celula = ws.cell(linha, coluna, valor)
                celula.font = Font(name=FONTE, size=14 if coluna == 5 else 12, bold=True)
                celula.alignment = ESQUERDA if coluna <= 2 else CENTRO
                celula.border = BORDA

            textos = [o["texto"] for o in pessoa["ocorrencias"]]
            celula = ws.cell(linha, 7, "\n".join(textos))
            celula.font = Font(name=FONTE, size=11, bold=bool(pessoa["sem_justificativa"]))
            celula.alignment = ESQUERDA
            celula.border = BORDA
            if pessoa["sem_justificativa"]:
                celula.fill = _preencher(AMARELO_PENDENTE)

            altura_texto = sum(max(1, -(-len(t) // 60)) for t in textos) if textos else 1
            ws.row_dimensions[linha].height = max(
                _altura_da_linha(pessoa["nome"], valores[1]), min(18 + altura_texto * 16, 240))
            self.linha += 1

        esquecidos = ponto.get("dias_sem_ponto") or []
        if esquecidos:
            linha = self.linha
            dias = ", ".join(f"{d['dia_semana'][:3]} {data_br(d['data'])[:5]}" for d in esquecidos)
            ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=5)
            celula = ws.cell(linha, 1, f"Ponto não lançado em: {dias}")
            celula.font = Font(name=FONTE, size=12, bold=True)
            celula.fill = _preencher(AMARELO_PENDENTE)
            celula.alignment = ESQUERDA
            for coluna in range(1, 6):
                ws.cell(linha, coluna).border = BORDA
            ws.row_dimensions[linha].height = 22
            self.linha += 1

        ws.row_dimensions[self.linha].height = 10.05
        self.linha += 1

    def secao(self, titulo: str, itens: list[dict], cor: str = CINZA_CABECALHO,
              entra_no_total: bool = True) -> None:
        ws = self.ws
        linha_titulo = self.linha

        ws.merge_cells(start_row=linha_titulo, start_column=1, end_row=linha_titulo, end_column=4)
        celula = ws.cell(linha_titulo, 1, titulo)
        celula.font = Font(name=FONTE, size=15, bold=True)
        celula.fill = _preencher(cor)
        celula.alignment = Alignment(horizontal="center", vertical="center")
        for coluna in range(1, 6):
            ws.cell(linha_titulo, coluna).border = BORDA
        ws.row_dimensions[linha_titulo].height = 34.95

        total = ws.cell(linha_titulo, 5)
        total.font = Font(name=FONTE, size=16, bold=True)
        total.fill = _preencher(cor)
        total.alignment = CENTRO
        total.number_format = MOEDA
        self.linha += 1

        primeira = self.linha
        for item in itens:
            self._linha_de_item(item)

        ultima = self.linha - 1
        if ultima >= primeira:
            total.value = f"=SUM(E{primeira}:E{ultima})"
        else:
            total.value = 0

        if entra_no_total:
            self.celulas_de_secao.append(f"E{linha_titulo}")

        ws.row_dimensions[self.linha].height = 10.05
        self.linha += 1

    def _linha_de_item(self, item: dict) -> None:
        ws = self.ws
        linha = self.linha

        fornecedor = item.get("fornecedor") or "-"
        descricao = item.get("descricao") or ""
        documento = item.get("documento") or "-"

        celula = ws.cell(linha, 1, fornecedor)
        celula.font = Font(name=FONTE, size=12, bold=True)
        celula.alignment = ESQUERDA

        celula = ws.cell(linha, 2, descricao)
        celula.font = Font(name=FONTE, size=12, bold=True)
        celula.alignment = ESQUERDA

        celula = ws.cell(linha, 3)
        data = item.get("data")
        if data:
            try:
                celula.value = datetime.strptime(str(data)[:10], "%Y-%m-%d")
                celula.number_format = DATA_BR
            except ValueError:
                celula.value = str(data)
        else:
            celula.value = "-"
        celula.font = Font(name=FONTE, size=12, bold=True)
        celula.alignment = CENTRO

        celula = ws.cell(linha, 4, documento)
        celula.font = Font(name=FONTE, size=12, bold=True)
        celula.alignment = CENTRO

        # Protegido com _safe_float
        valor = round(_safe_float(item.get("valor")), 2)
        celula = ws.cell(linha, 5, valor)
        celula.font = Font(name=FONTE, size=14, bold=True)
        celula.alignment = CENTRO
        celula.number_format = MOEDA

        for coluna in range(1, 6):
            ws.cell(linha, coluna).border = BORDA

        if item.get("status"):
            marca = ws.cell(linha, 6, item["status"])
            marca.font = Font(name=FONTE, size=14)
            marca.fill = _preencher(CINZA_STATUS)
            marca.alignment = Alignment(vertical="center")

        ws.row_dimensions[linha].height = _altura_da_linha(fornecedor, descricao, str(documento))
        self.linha += 1

    def total_geral(self) -> None:
        ws = self.ws
        self.linha += 1
        linha = self.linha

        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=3)
        celula = ws.cell(linha, 1, f"TOTAL A PAGAR - SEMANA {self.semana:02d}")
        celula.font = Font(name=FONTE, size=20, bold=True, color="FFFFFFFF")
        celula.fill = _preencher(CINZA_CABECALHO)
        celula.alignment = CENTRO

        ws.merge_cells(start_row=linha, start_column=4, end_row=linha, end_column=5)
        total = ws.cell(linha, 4)
        total.value = ("=" + "+".join(self.celulas_de_secao)) if self.celulas_de_secao else 0
        total.font = Font(name=FONTE, size=20, bold=True, color="FFFFFFFF")
        total.fill = _preencher(CINZA_CABECALHO)
        total.alignment = CENTRO
        total.number_format = MOEDA

        for coluna in range(1, 6):
            ws.cell(linha, coluna).border = BORDA
        ws.row_dimensions[linha].height = 40

        ws.print_area = f"A1:{self.ultima_coluna}{linha}"


# ---------------------------------------------------------------- dados

ROTULO_STATUS = {"Em Aberto": "A PAGAR", "Reembolso": "REEMBOLSO", "Pago": "PAGO"}


def _itens_de_despesas(despesas: list[dict], status_alvo: str) -> list[dict]:
    return [{
        "fornecedor": d.get("fornecedor") or d.get("categoria") or "-",
        "descricao": (d.get("descricao") or "") + (
            f" · reembolsar {d['reembolsar_para']}" if d.get("reembolsar_para") else ""),
        "data": d.get("data"),
        "documento": d.get("numero_nota") or "-",
        "valor": d.get("valor") or 0,
        "status": ROTULO_STATUS.get(d.get("status_pagamento"), ""),
    } for d in despesas
      if d.get("status_pagamento") == status_alvo]


def _horas_br(valor) -> str:
    return f"{round(_safe_float(valor), 2):g}".replace(".", ",") + " h"


def montar_aba(planilha, obra: dict, dados: dict) -> None:
    ponto = dados.get("ponto_equipe") or {}
    tem_ponto = bool(ponto.get("equipe") or ponto.get("dias_sem_ponto"))
    aba = _Aba(planilha, obra, dados, com_ocorrencias=tem_ponto)
    aba.cabecalho()

    despesas = dados.get("despesas") or []
    semana = dados.get("semana_numero", 0)

    aba.secao(f" NOTAS PARA REEMBOLSO SEMANA {semana:02d}",
              _itens_de_despesas(despesas, "Reembolso"))
    aba.secao(f"NOTAS ABERTAS A PAGAR {aba.sigla}",
              _itens_de_despesas(despesas, "Em Aberto"))

    # A antiga "folha" (horas × valor/hora) saiu: o relatório mostra só a
    # presença. Pagamento da equipe entra no Caixa como despesa.
    if tem_ponto:
        aba.secao_ponto(f"PONTO DA EQUIPE - SEMANA {semana:02d}", ponto)

    pagas = _itens_de_despesas(despesas, "Pago")
    if pagas:
        aba.secao(f"NOTAS JÁ PAGAS NA SEMANA {semana:02d} (fora do total)", pagas,
                  entra_no_total=False)

    aba.total_geral()


def _nome_da_aba(obra: dict, usados: set[str]) -> str:
    base = f"RelatorioLancamento {_sigla_da_obra(obra)}"[:31]
    nome, contador = base, 2
    while nome in usados:
        sufixo = f" {contador}"
        nome = base[:31 - len(sufixo)] + sufixo
        contador += 1
    usados.add(nome)
    return nome


def gerar_planilha(abas: list[tuple[dict, dict]]) -> bytes:
    planilha = Workbook()
    planilha.remove(planilha.active)

    usados: set[str] = set()
    for obra, dados in abas:
        aba = planilha.create_sheet(_nome_da_aba(obra, usados))
        montar_aba(aba, obra, dados)

    buffer = io.BytesIO()
    planilha.save(buffer)
    return buffer.getvalue()


def nome_arquivo(obra: dict, dados: dict) -> str:
    def limpo(texto: str) -> str:
        return re.sub(r"[^A-Za-z0-9]+", "_", (texto or "").upper()).strip("_")

    inicio = data_br(dados.get("data_inicio") or "INICIO").replace("/", "_")
    fim = data_br(dados.get("data_fim") or "FIM").replace("/", "_")
    return f"REL_DE_PAGAMENTO_{limpo(_sigla_da_obra(obra))}_{inicio}_A_{fim}.xlsx"
