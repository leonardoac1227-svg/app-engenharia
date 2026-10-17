# -*- coding: utf-8 -*-
"""Faltas no ponto (com justificativa) e relatório de pagamento sem valor por funcionário."""
import io
import os
import sys
import tempfile

# Caminho do projeto derivado da posição deste arquivo: assim a suíte roda em
# qualquer clone, sem depender de onde a pasta foi colocada.
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
sys.path.insert(0, os.path.join(RAIZ, "ferramentas"))

os.environ.setdefault("DATA_DIR", tempfile.mkdtemp(prefix="obra-teste-"))
from openpyxl import load_workbook

from obra import create_app
from obra.db import executar

falhas = []
def ok(cond, msg, extra=''):
    print(('  OK  ' if cond else ' FALHA') + ' ' + msg + ('' if cond else f'  -> {extra}'))
    if not cond: falhas.append(msg)

app = create_app(); c = app.test_client()
c.post('/api/auth/login', json={'email':'gestor@obra.local','senha':'obra123'})
H = {'X-CSRF-Token': c.get('/api/sessao').get_json()['csrf']}
O = c.post('/api/obras', json={'nome':'Obra Ponto 22'}, headers=H).get_json()['id']
S = c.post(f'/api/obras/{O}/semanas', json={'data_inicio':'2026-08-17'}, headers=H).get_json()['id']
dias = c.get(f'/api/semanas/{S}/dias').get_json()
SEG, TER, QUA, QUI, SEX, SAB = (d['id'] for d in dias[:6])

def func(nome):
    return c.post(f'/api/obras/{O}/funcionarios', json={'nome': nome, 'funcao': 'Pedreiro'},
                  headers=H).get_json()['id']
JOAO, MARIA, PEDRO, NOVATO = func('JOÃO'), func('MARIA'), func('PEDRO'), func('NOVATO')
with app.app_context():
    # Todos já estavam na obra na semana testada, menos o novato (entrou na quinta).
    executar("UPDATE funcionarios SET criado_em = '2026-08-01T08:00:00' WHERE obra_id = ?", (O,))
    executar("UPDATE funcionarios SET criado_em = '2026-08-20T08:00:00' WHERE id = ?", (NOVATO,))

def ponto(dia, registros):
    return c.put(f'/api/dias/{dia}/horas', json={'registros': [
        {'funcionario_id': f, 'horas': h, 'justificativa': j} for f, h, j in registros]}, headers=H)

print('== JUSTIFICATIVA NO PONTO ==')
ponto(SEG, [(JOAO, 8, 'não deveria ficar'), (MARIA, 4, 'Atestado médico'), (PEDRO, 0, 'Chuva / obra parada')])
h = c.get(f'/api/dias/{SEG}/horas').get_json()
ok(h[str(JOAO)]['justificativa'] == '', 'jornada cumprida não guarda justificativa', h[str(JOAO)])
ok(h[str(MARIA)] == {'horas': 4, 'justificativa': 'Atestado médico'}, 'meio turno guarda o motivo', h[str(MARIA)])
ok(h[str(PEDRO)] == {'horas': 0, 'justificativa': 'Chuva / obra parada'},
   'falta justificada fica registrada com 0 h', h.get(str(PEDRO)))

ponto(TER, [(JOAO, 7, ''), (MARIA, 8, ''), (PEDRO, 0, '')])
h = c.get(f'/api/dias/{TER}/horas').get_json()
ok(str(PEDRO) not in h, 'falta sem justificativa não cria registro (como antes)', h)
ponto(QUA, [(JOAO, 8, ''), (MARIA, 8, ''), (PEDRO, 3, 'Saiu mais cedo (combinado)')])
ponto(QUI, [(JOAO, 8, ''), (MARIA, 8, ''), (PEDRO, 8, '')])          # novato sem ponto: falta
# sexta: ninguém lançou ponto → dia esquecido, não falta de todo mundo
ponto(SAB, [(JOAO, 4, '')])                                          # sábado é extra

print('== OCORRÊNCIAS DA SEMANA ==')
fe = c.get(f'/api/semanas/{S}/fechamento').get_json()
pe = fe['ponto_equipe']
pessoa = {p['nome']: p for p in pe['equipe']}
tipos = lambda nome: [(o['dia_semana'], o['tipo']) for o in pessoa[nome]['ocorrencias']]

ok(pe['jornada'] == 8, 'jornada padrão de 8 h', pe['jornada'])
ok(tipos('JOÃO') == [('Terça', 'parcial')], 'João: só a terça (7 h); sábado não conta', tipos('JOÃO'))
ok(pessoa['JOÃO']['horas'] == 35 and pessoa['JOÃO']['dias_trabalhados'] == 5,
   'horas da semana incluem o sábado', pessoa['JOÃO'])
ok(tipos('MARIA') == [('Segunda', 'meio_turno')], 'Maria: meio turno na segunda', tipos('MARIA'))
ok(tipos('PEDRO') == [('Segunda', 'falta'), ('Terça', 'falta'), ('Quarta', 'parcial')],
   'Pedro: falta, falta e 3 h', tipos('PEDRO'))
ok(tipos('NOVATO') == [('Quinta', 'falta')], 'novato só é cobrado a partir do cadastro', tipos('NOVATO'))
ok(pessoa['PEDRO']['sem_justificativa'] == 1, 'conta faltas sem justificativa', pessoa['PEDRO'])
ok([d['dia_semana'] for d in pe['dias_sem_ponto']] == ['Sexta'], 'sexta aparece como ponto não lançado',
   pe['dias_sem_ponto'])

textos = {o['tipo'] + ':' + o['dia_semana']: o['texto'] for p in pe['equipe'] for o in p['ocorrencias']}
ok(textos['meio_turno:Segunda'] == 'Seg 17/08: falta de meio turno (4 h) · Atestado médico',
   'texto do meio turno', textos['meio_turno:Segunda'])
ok(textos['parcial:Terça'] == 'Ter 18/08: 7 h de 8 h (faltaram 1 h) · sem justificativa',
   'texto de horas quebradas', textos['parcial:Terça'])
ok('Qua 19/08: 3 h de 8 h (faltaram 5 h) · Saiu mais cedo (combinado)' in pessoa['PEDRO']['ocorrencias'][2]['texto'],
   'texto com justificativa', pessoa['PEDRO']['ocorrencias'][2]['texto'])

print('== RELATÓRIO DE PAGAMENTO SEM PREÇO ==')
# Ponto antigo com valor/hora gravado: era isso que fazia o preço aparecer.
with app.app_context():
    executar("UPDATE horas SET valor_hora = 30 WHERE dia_id = ?", (QUI,))
c.post(f'/api/dias/{SEG}/despesas', data={'categoria':'Material','descricao':'Cimento','valor':'100',
       'status_pagamento':'Em Aberto'}, headers=H, content_type='multipart/form-data')

r = c.get(f'/api/semanas/{S}/pagamentos.xlsx')
ok(r.status_code == 200, 'planilha sai', r.status_code)
ws = load_workbook(io.BytesIO(r.data)).active
celulas = [[ws.cell(l, col).value for col in range(1, 8)] for l in range(1, ws.max_row + 1)]
texto_todo = '\n'.join(str(v) for linha in celulas for v in linha if v is not None)
ok('FOLHA' not in texto_todo.upper().replace('FOLHA DE PONTO', ''), 'seção de folha com valor sumiu')
ok('PONTO DA EQUIPE - SEMANA 01' in texto_todo, 'seção de ponto da equipe aparece')
linha_joao = next(l for l in celulas if l[0] == 'JOÃO')
ok(linha_joao[4] == '35 h' and linha_joao[3] == 'PONTO', 'linha do funcionário mostra horas, não R$', linha_joao)
ok(not any(isinstance(l[4], (int, float)) and l[4] > 0 and l[3] == 'PONTO' for l in celulas),
   'nenhum valor em dinheiro nas linhas do ponto')
ok(linha_joao[6] == 'Ter 18/08: 7 h de 8 h (faltaram 1 h) · sem justificativa',
   'coluna lateral com a ocorrência do dia', linha_joao[6])
linha_maria = next(l for l in celulas if l[0] == 'MARIA')
ok('meio turno' in linha_maria[6] and 'Atestado médico' in linha_maria[6], 'meio turno com motivo', linha_maria[6])
linha_pedro = next(l for l in celulas if l[0] == 'PEDRO')
ok(linha_pedro[6].count('\n') == 2, 'várias ocorrências, uma por linha', linha_pedro[6])
ok(any(str(l[0] or '').startswith('Ponto não lançado em: Sex') for l in celulas), 'avisa o dia sem ponto')

total = next(l for l in celulas if str(l[0] or '').startswith('TOTAL A PAGAR'))
somadas = [celulas[int(ref[1:]) - 1][0] for ref in total[3][1:].split('+')]
ok(len(somadas) == 2 and 'REEMBOLSO' in somadas[0] and 'ABERTAS' in somadas[1],
   'total a pagar soma só reembolso e notas abertas (sem folha)', (total[3], somadas))
ok(next(l for l in celulas if l[0] == 'Fornecedor')[6] == 'Ocorrências do ponto', 'cabeçalho da coluna lateral')

pdf = c.get(f'/api/semanas/{S}/relatorio.pdf')
ok(pdf.status_code == 200 and pdf.data[:4] == b'%PDF', 'PDF continua saindo')
txt = c.get(f'/api/semanas/{S}/relatorio.txt').get_data(as_text=True)
ok('PONTO DA EQUIPE' in txt and 'Atestado médico' in txt and 'FOLHA POR PROFISSIONAL' not in txt,
   'relatório em texto traz o ponto sem valores')
ok(c.get('/api/inicio').get_json().get('jornada_horas') == 8, 'app recebe a jornada no início')

print('== REPETIR ÚLTIMO DIA ==')
c.post(f'/api/dias/{SEX}/horas/repetir', json={}, headers=H)
h = c.get(f'/api/dias/{SEX}/horas').get_json()
ok(str(PEDRO) not in h or h[str(PEDRO)]['justificativa'] == '',
   'repetir o ponto não copia falta nem justificativa', h)

print('\n' + ('PONTO OK' if not falhas else f'{len(falhas)} FALHA(S): {falhas}'))
sys.exit(1 if falhas else 0)
